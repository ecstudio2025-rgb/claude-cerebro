#!/usr/bin/env python3
"""
ECSSTUDIO — Viral Profile Scraper
Niche-agnostic TikTok scraper que extrae guiones de videos virales.
Uso: python3 viral_scraper.py --config config.json
"""

import argparse, json, os, re, sys, time
from datetime import datetime, timezone, timedelta
from pathlib import Path

try:
    import requests
    from deep_translator import GoogleTranslator
    from langdetect import detect, LangDetectException
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError as e:
    print(f"[ERROR] Dependencia faltante: {e}")
    print("Instalar con: pip install requests deep-translator langdetect openpyxl --break-system-packages")
    sys.exit(1)

# ─── APIFY ────────────────────────────────────────────────────────────────────

APIFY_BASE = "https://api.apify.com/v2"
TT_ACTOR   = "clockworks~free-tiktok-scraper"
POLL_SECS  = 10

def log(msg: str) -> None:
    print(f"[{datetime.now().strftime('%H:%M:%S')}] {msg}", flush=True)

def apify_start_run(token: str, hashtag: str, results: int) -> str | None:
    """Inicia un run de TikTok scraper. Devuelve run_id o None si falla."""
    try:
        r = requests.post(
            f"{APIFY_BASE}/acts/{TT_ACTOR}/runs?token={token}",
            json={"hashtags": [f"#{hashtag}"], "resultsPerPage": results, "maxResultsPerQuery": results},
            timeout=30
        )
        if r.status_code == 402:
            log(f"  ✗ #{hashtag} — Sin créditos Apify (402). Recargar cuenta en console.apify.com")
            return None
        r.raise_for_status()
        return r.json()["data"]["id"]
    except Exception as e:
        log(f"  ✗ #{hashtag} — Error al iniciar: {e}")
        return None

def apify_poll(token: str, run_id: str) -> str:
    """Consulta el estado de un run. Devuelve status string."""
    try:
        r = requests.get(f"{APIFY_BASE}/actor-runs/{run_id}?token={token}", timeout=15)
        return r.json()["data"]["status"]
    except:
        return "UNKNOWN"

def apify_get_items(token: str, run_id: str) -> list:
    """Obtiene los items del dataset de un run completado."""
    try:
        r = requests.get(
            f"{APIFY_BASE}/actor-runs/{run_id}/dataset/items?token={token}&limit=500&clean=true",
            timeout=30
        )
        data = r.json()
        return data if isinstance(data, list) else []
    except:
        return []

# ─── TRANSCRIPT ───────────────────────────────────────────────────────────────

def parse_vtt(content: str) -> str:
    """Parsea un archivo WebVTT y devuelve el texto limpio sin timestamps."""
    texts = []
    for line in content.strip().splitlines():
        line = line.strip()
        if not line or line == "WEBVTT" or "-->" in line or re.match(r'^\d+$', line):
            continue
        line = re.sub(r'<[^>]+>', '', line)
        line = re.sub(r'\{\\[^}]+\}', '', line)
        if line:
            texts.append(line)
    # Deduplicar líneas consecutivas idénticas
    deduped = []
    for t in texts:
        if not deduped or t != deduped[-1]:
            deduped.append(t)
    return ' '.join(deduped).strip()

def get_transcript(subtitle_links: list) -> tuple[str | None, str | None]:
    """Descarga el mejor subtítulo disponible. Devuelve (transcript, lang_code)."""
    if not subtitle_links:
        return None, None
    # Prioridad: ASR > LC > MT
    def priority(sl):
        s = sl.get("source", "")
        return 0 if s == "ASR" else (1 if s == "LC" else 2)
    for sl in sorted(subtitle_links, key=priority)[:4]:
        try:
            r = requests.get(
                sl["downloadLink"],
                headers={"User-Agent": "Mozilla/5.0", "Referer": "https://www.tiktok.com/"},
                timeout=12
            )
            if r.status_code == 200:
                text = parse_vtt(r.text)
                if text and len(text) > 30:
                    return text, sl.get("language", "?")
        except:
            pass
    return None, None

# ─── LANGUAGE & TRANSLATION ───────────────────────────────────────────────────

def detect_lang(text: str) -> str:
    """Detecta el idioma de un texto. Devuelve código ISO (es, en, pt, ru…)."""
    try:
        return detect(text[:600]) if len(text) > 30 else "unknown"
    except LangDetectException:
        return "unknown"

def translate(text: str, target_lang: str) -> str:
    """Traduce texto al idioma objetivo. Chunkeado para textos largos."""
    if not text or len(text) < 5:
        return text or ""
    try:
        chunks = [text[i:i+4500] for i in range(0, len(text), 4500)]
        result = []
        for ch in chunks:
            t = GoogleTranslator(source="auto", target=target_lang).translate(ch)
            if t:
                result.append(t)
            if len(chunks) > 1:
                time.sleep(0.4)
        return " ".join(result) if result else text
    except Exception as e:
        log(f"  ⚠ Error traducción: {e}")
        return text

def is_target_lang(detected: str, target: str) -> bool:
    """Comprueba si el idioma detectado coincide con el objetivo."""
    d = (detected or "").lower()
    t = (target or "es").lower()
    lang_map = {
        "es": ["es", "spa", "ca"],
        "en": ["en", "eng"],
        "pt": ["pt", "por"],
        "ru": ["ru", "rus"],
        "fr": ["fr", "fre", "fra"],
        "de": ["de", "deu", "ger"],
        "it": ["it", "ita"],
    }
    prefixes = lang_map.get(t, [t])
    return any(d.startswith(p) for p in prefixes)

# ─── PROCESSING ───────────────────────────────────────────────────────────────

def ts_to_dt(ts) -> datetime | None:
    """Convierte un timestamp Unix a datetime UTC."""
    try:
        return datetime.fromtimestamp(int(ts), tz=timezone.utc)
    except:
        return None

def fmt_views(n) -> str:
    """Formatea un número de views (1.2M, 450K, etc.)."""
    try:
        n = int(n)
        if n >= 1_000_000: return f"{n/1_000_000:.1f}M"
        if n >= 1_000:     return f"{n/1_000:.0f}K"
        return str(n)
    except:
        return "—"

def process_item(item: dict, market: dict, output_language: str, cutoff: datetime) -> dict | None:
    """Procesa un item de TikTok. Devuelve dict enriquecido o None si hay que descartarlo."""
    views = item.get("playCount", 0)
    ct    = item.get("createTime", 0)
    dt    = ts_to_dt(ct)
    if not dt or dt < cutoff:
        return None

    author  = item.get("authorMeta", {})
    handle  = author.get("name") or author.get("uniqueId") or ""
    caption = (item.get("text") or "").strip()

    # Obtener transcript
    sub_links   = (item.get("videoMeta") or {}).get("subtitleLinks") or []
    trans_orig, lang_sub = get_transcript(sub_links)

    # Fallback a caption si no hay subtítulos o el texto es muy corto
    if not trans_orig or len(trans_orig) < 40:
        trans_orig = caption
        lang_sub   = item.get("textLanguage") or market.get("lang_code") or "?"

    # Descartar si sigue siendo basura (solo hashtags, sin contenido hablado)
    word_count = len([w for w in trans_orig.split() if not w.startswith("#")])
    if word_count < 5:
        return None

    # Detectar idioma real del transcript
    lang_detected = detect_lang(trans_orig)

    # Traducir si es necesario
    already_target = is_target_lang(lang_detected, output_language)
    if already_target:
        trans_out = trans_orig
        translated = False
    else:
        trans_out  = translate(trans_orig, output_language)
        translated = True

    return {
        "uid":          f"{market['id']}_{handle}_{item.get('id', '')}",
        "market_id":    market["id"],
        "market_name":  market["name"],
        "market_flag":  market.get("flag", ""),
        "market_color": market.get("color", "9E9E9E"),
        "handle":       f"@{handle}",
        "display_name": author.get("nickName") or handle,
        "followers":    author.get("fans") or "—",
        "views":        views,
        "views_fmt":    fmt_views(views),
        "likes":        item.get("diggCount", 0),
        "comments":     item.get("commentCount", 0),
        "date":         dt.strftime("%Y-%m-%d"),
        "days_ago":     (datetime.now(tz=timezone.utc) - dt).days,
        "lang_orig":    lang_detected or lang_sub or "?",
        "lang_sub":     lang_sub or "?",
        "caption":      caption,
        "transcript":   trans_orig,
        "transcript_out": trans_out,
        "translated":   translated,
        "url":          item.get("webVideoUrl") or "",
        "hashtag":      item.get("searchHashtag") or "",
    }

def extract_hook(text: str) -> str:
    """Extrae la primera frase significativa de un transcript."""
    if not text:
        return ""
    sentences = re.split(r'(?<=[.!?¿¡])\s+', text)
    hook = next((s.strip() for s in sentences if len(s.strip()) >= 20), text[:200].strip())
    return re.sub(r'\s+', ' ', hook)[:250]

# ─── XLSX ─────────────────────────────────────────────────────────────────────

def build_xlsx(results: list, config: dict, now: datetime) -> str:
    """Genera el XLSX con 4 hojas. Devuelve la ruta del archivo."""

    # ── Estilos base ──────────────────────────────────────────────────────────
    W       = "FFFFFF"
    G1      = "F7F9FF"
    G2      = "EFF2FA"
    DARK    = "1E2A3A"
    HDR2    = "2D3E50"
    MUTED   = "8899AA"
    TEXT    = "2C2C3E"
    GOLD    = "C98B00"
    GREEN   = "1A7850"
    TRANS_BG = "F0FBF7"
    BORDER  = "D8DCF0"

    def F(hex_c): return PatternFill(start_color=hex_c, end_color=hex_c, fill_type="solid")
    def fn(bold=False, size=9, color=TEXT, underline=None):
        kw = dict(bold=bold, size=size, color=color, name="Arial")
        if underline: kw["underline"] = underline
        return Font(**kw)
    def brd():
        t = Side(style="thin", color=BORDER)
        return Border(left=t, right=t, top=t, bottom=t)
    def al(h="left", v="top", wrap=False):
        return Alignment(horizontal=h, vertical=v, wrap_text=wrap, indent=1 if h=="left" else 0)

    bg_alt = [F(W), F(G1)]

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    client_name = config.get("client_name", "Cliente")
    niche       = config.get("client_niche", "")
    out_lang    = config.get("output_language", "es").upper()

    def banner(ws, row, text, bg=DARK, fg="FFFFFF", size=10, height=24):
        ws.merge_cells(f"A{row}:{get_column_letter(ws.max_column or 14)}{row}")
        c = ws.cell(row=row, column=1)
        c.value = text
        c.font = Font(bold=True, size=size, color=fg, name="Arial")
        c.fill = F(bg)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=2)
        ws.row_dimensions[row].height = height

    # ── SHEET 1: Guiones Virales ───────────────────────────────────────────────
    ws = wb.create_sheet("Guiones Virales")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    col_defs = [
        ("#",           3,  "center"),
        ("Mercado",    14,  "left"),
        ("Cuenta",     20,  "left"),
        ("Nombre",     22,  "left"),
        ("Seguidores", 13,  "center"),
        ("Views",      11,  "center"),
        ("Likes",       9,  "center"),
        ("Fecha",      12,  "center"),
        ("Días atrás",  9,  "center"),
        ("Idioma orig", 13, "center"),
        ("Caption",    30,  "left"),
        ("Guion original", 55, "left"),
        (f"Guion [{out_lang}]", 55, "left"),
        ("URL",        35,  "left"),
    ]
    for ci, (label, width, _) in enumerate(col_defs, 1):
        ws.column_dimensions[get_column_letter(ci)].width = width

    # Row 1 title
    ws.merge_cells(f"A1:{get_column_letter(len(col_defs))}1")
    c = ws["A1"]
    c.value = f"GUIONES VIRALES · {client_name.upper()} · {niche} · {now.strftime('%d/%m/%Y')}"
    c.font = Font(bold=True, size=10, color="FFFFFF", name="Arial")
    c.fill = F(DARK)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=2)
    ws.row_dimensions[1].height = 24

    # Row 2 headers
    for ci, (label, _, align) in enumerate(col_defs, 1):
        c = ws.cell(row=2, column=ci)
        c.value = label
        c.font = Font(bold=True, size=8, color="FFFFFF", name="Arial")
        c.fill = F(HDR2)
        c.alignment = al(align, "center")
        c.border = brd()
    ws.row_dimensions[2].height = 20

    for ri, r in enumerate(results):
        row = ri + 3
        bg  = bg_alt[ri % 2]
        mkt_col = r["market_color"]
        vnum    = r["views"]
        vcol    = "8B0000" if vnum >= 5_000_000 else (GOLD if vnum >= 1_000_000 else GREEN)

        def W_cell(col, val, bold=False, color=TEXT, h="left", wrap=False, ufill=None):
            c = ws.cell(row=row, column=col)
            c.value = val
            c.font  = fn(bold=bold, color=color)
            c.fill  = ufill or bg
            c.alignment = al(h, "top", wrap)
            c.border = brd()

        W_cell(1, ri+1, h="center", color=MUTED)
        # Market badge
        c_mkt = ws.cell(row=row, column=2)
        c_mkt.value = f"{r['market_flag']} {r['market_id']}"
        c_mkt.font  = Font(bold=True, size=9, color=mkt_col, name="Arial")
        c_mkt.fill  = F(G2)
        c_mkt.alignment = al("center", "top")
        c_mkt.border = brd()

        W_cell(3, r["handle"],       bold=True, color="1565C0")
        W_cell(4, r["display_name"])
        W_cell(5, r["followers"],    h="center", color=MUTED)
        # Views
        c_v = ws.cell(row=row, column=6)
        c_v.value = r["views_fmt"]
        c_v.font  = Font(bold=True, size=10 if vnum>=4_000_000 else 9, color=vcol, name="Arial")
        c_v.fill  = bg
        c_v.alignment = al("center", "top")
        c_v.border = brd()

        W_cell(7, fmt_views(r["likes"]),    h="center", color=MUTED)
        W_cell(8, r["date"],                h="center", color="555566")
        W_cell(9, r["days_ago"],            h="center", color=MUTED)
        W_cell(10, r["lang_orig"],          h="center", color=mkt_col, ufill=F(G2))
        W_cell(11, r["caption"],            color="888899", wrap=True)
        W_cell(12, r["transcript"],         wrap=True)

        # Translated column
        es_bg  = F(TRANS_BG) if r["translated"] else bg
        es_col = "1A3A2A"    if r["translated"] else TEXT
        c_es   = ws.cell(row=row, column=13)
        c_es.value = r["transcript_out"]
        c_es.font  = fn(color=es_col)
        c_es.fill  = es_bg
        c_es.alignment = al("left", "top", True)
        c_es.border = brd()

        # URL
        c_url = ws.cell(row=row, column=14)
        c_url.value = r["url"]
        if r["url"].startswith("http"):
            c_url.hyperlink = r["url"]
        c_url.font  = fn(color="1565C0", underline="single", size=8)
        c_url.fill  = bg
        c_url.alignment = al("left", "top")
        c_url.border = brd()

        ws.row_dimensions[row].height = max(50, min(200, len(r["transcript_out"]) // 4 + 20))

    ws.sheet_properties.tabColor = "1E2A3A"

    # ── SHEET 2: Ganchos ──────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Ganchos")
    ws2.sheet_view.showGridLines = False
    ws2.freeze_panes = "A3"

    ws2.merge_cells("A1:E1")
    c = ws2["A1"]
    c.value = f"GANCHOS · Primeras frases de cada guion — extraer patrones de hook"
    c.font = Font(bold=True, size=10, color="FFFFFF", name="Arial")
    c.fill = F(DARK)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=2)
    ws2.row_dimensions[1].height = 22

    hdr2 = [("Mercado",14),("Cuenta",20),("Views",11),("Días atrás",10),("Hook (primera frase del guion)",110)]
    for ci, (h, w) in enumerate(hdr2, 1):
        c = ws2.cell(row=2, column=ci)
        c.value = h
        c.font  = Font(bold=True, size=8, color="FFFFFF", name="Arial")
        c.fill  = F(HDR2)
        c.alignment = al("center" if ci<5 else "left", "center")
        c.border = brd()
        ws2.column_dimensions[get_column_letter(ci)].width = w
    ws2.row_dimensions[2].height = 20

    for ri, r in enumerate(results):
        row = ri + 3
        bg  = bg_alt[ri % 2]
        hook = extract_hook(r["transcript_out"])
        vnum = r["views"]
        vcol = "8B0000" if vnum>=5_000_000 else (GOLD if vnum>=1_000_000 else GREEN)
        mkt_col = r["market_color"]

        data = [
            (f"{r['market_flag']} {r['market_id']}", "center", mkt_col, True,  F(G2)),
            (r["handle"],                             "center", "1565C0",True,  bg),
            (r["views_fmt"],                          "center", vcol,    True,  bg),
            (r["days_ago"],                           "center", MUTED,   False, bg),
            (hook,                                    "left",   TEXT,    False, F(TRANS_BG) if r["translated"] else bg),
        ]
        for ci, (val, h, col, bold, cbg) in enumerate(data, 1):
            c = ws2.cell(row=row, column=ci)
            c.value = val
            c.font  = fn(bold=bold, color=col)
            c.fill  = cbg
            c.alignment = al(h, "top", ci==5)
            c.border = brd()
        ws2.row_dimensions[row].height = max(30, min(80, len(hook)//3))

    ws2.sheet_properties.tabColor = "2D5F8A"

    # ── SHEET 3: Dashboard ─────────────────────────────────────────────────────
    ws3 = wb.create_sheet("Dashboard")
    ws3.sheet_view.showGridLines = False

    ws3.merge_cells("A1:F1")
    c = ws3["A1"]
    c.value = f"DASHBOARD · {client_name} · {now.strftime('%d/%m/%Y')}"
    c.font  = Font(bold=True, size=12, color="FFFFFF", name="Arial")
    c.fill  = F(DARK)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=2)
    ws3.row_dimensions[1].height = 28

    # KPIs
    total_views = sum(r["views"] for r in results)
    mkt_counts  = {}
    for r in results:
        mkt_counts[r["market_id"]] = mkt_counts.get(r["market_id"], 0) + 1

    kpis = [
        ("Videos encontrados",  len(results),                      "1565C0"),
        ("Views totales",        fmt_views(total_views),            GOLD),
        ("Mercados rastreados",  len(mkt_counts),                   GREEN),
        ("Con transcripción",    sum(1 for r in results if len(r["transcript"])>100), "6B3FA0"),
        ("Traducidos",           sum(1 for r in results if r["translated"]),          "0D7A5A"),
        ("Días de ventana",      config.get("date_range_days",90),  MUTED),
    ]
    for ci, (col_w) in enumerate([18]*6, 1):
        ws3.column_dimensions[get_column_letter(ci)].width = 18

    for ci, (label, val, color) in enumerate(kpis, 1):
        cv = ws3.cell(row=3, column=ci)
        cv.value = val
        cv.font  = Font(bold=True, size=20, color=color, name="Arial")
        cv.fill  = F(W)
        cv.alignment = Alignment(horizontal="center", vertical="center")
        cv.border = brd()
        cl = ws3.cell(row=4, column=ci)
        cl.value = label
        cl.font  = Font(bold=True, size=9, color=TEXT, name="Arial")
        cl.fill  = F(G1)
        cl.alignment = Alignment(horizontal="center", vertical="center")
        cl.border = brd()
        ws3.row_dimensions[3].height = 38
        ws3.row_dimensions[4].height = 24

    # Top 10
    ws3.row_dimensions[6].height = 22
    ws3.merge_cells("A6:F6")
    c = ws3["A6"]
    c.value = "TOP 10 VIDEOS POR VIEWS"
    c.font  = Font(bold=True, size=10, color="FFFFFF", name="Arial")
    c.fill  = F(HDR2)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=2)
    c.border = brd()

    top10_hdrs = [("Mercado",14),("Cuenta",18),("Views",10),("Días",8),("Idioma",12),("Gancho",70)]
    for ci, (h, w) in enumerate(top10_hdrs, 1):
        c = ws3.cell(row=7, column=ci)
        c.value = h
        c.font  = Font(bold=True, size=8, color="FFFFFF", name="Arial")
        c.fill  = F(DARK)
        c.alignment = al("center" if ci<6 else "left", "center")
        c.border = brd()
        ws3.column_dimensions[get_column_letter(ci)].width = w
    ws3.row_dimensions[7].height = 18

    for i, r in enumerate(sorted(results, key=lambda x: x["views"], reverse=True)[:10]):
        row  = i + 8
        bg   = bg_alt[i % 2]
        vnum = r["views"]
        vcol = "8B0000" if vnum>=5_000_000 else (GOLD if vnum>=1_000_000 else GREEN)
        hook = extract_hook(r["transcript_out"])
        mkt_col = r["market_color"]
        data = [
            (f"{r['market_flag']} {r['market_id']}", "center", mkt_col, True,  F(G2)),
            (r["handle"],                             "center", "1565C0",True,  bg),
            (r["views_fmt"],                          "center", vcol,    True,  bg),
            (r["days_ago"],                           "center", MUTED,   False, bg),
            (r["lang_orig"],                          "center", mkt_col, False, F(G2)),
            (hook,                                    "left",   TEXT,    False, bg),
        ]
        for ci, (val, h, col, bold, cbg) in enumerate(data, 1):
            c = ws3.cell(row=row, column=ci)
            c.value = val
            c.font  = fn(bold=bold, color=col)
            c.fill  = cbg
            c.alignment = al(h, "center", ci==6)
            c.border = brd()
        ws3.row_dimensions[row].height = 22

    ws3.sheet_properties.tabColor = "2D7A4F"

    # ── SHEET 4: Metodología ───────────────────────────────────────────────────
    ws4 = wb.create_sheet("Metodología")
    ws4.sheet_view.showGridLines = False
    ws4.column_dimensions["A"].width = 35
    ws4.column_dimensions["B"].width = 60

    ws4.merge_cells("A1:B1")
    c = ws4["A1"]
    c.value = "PARÁMETROS DEL SCRAPING"
    c.font  = Font(bold=True, size=11, color="FFFFFF", name="Arial")
    c.fill  = F(DARK)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=2)
    ws4.row_dimensions[1].height = 26

    all_hashtags = []
    for m in config.get("markets", []):
        tags = ", ".join(f"#{t}" for t in m.get("hashtags", []))
        all_hashtags.append(f"{m.get('flag','')} {m['id']}: {tags}")

    stats = [
        ("Fecha extracción",         now.strftime("%d/%m/%Y %H:%M UTC")),
        ("Cliente",                  client_name),
        ("Nicho",                    niche),
        ("Actor Apify",              TT_ACTOR),
        ("Umbral de vistas",         f"{config.get('min_views',1_000_000):,}"),
        ("Ventana temporal",         f"{config.get('date_range_days',90)} días"),
        ("Resultados por hashtag",   config.get("results_per_hashtag", 100)),
        ("Idioma de salida",         out_lang),
        ("Videos encontrados",       len(results)),
        ("Mercados rastreados",      " · ".join(m["id"] for m in config.get("markets",[]))),
        ("Hashtags rastreados",      "\n".join(all_hashtags)),
        ("Fuente transcripciones",   "TikTok ASR nativo (WebVTT via subtitleLinks)"),
        ("Motor de traducción",      "GoogleTranslator via deep_translator (gratuito)"),
        ("Nota créditos Apify",      "Plan Free = $5/mes ≈ 8–12 runs. Para uso regular recomendado plan Starter ($49/mes)."),
    ]

    for i, (k, v) in enumerate(stats):
        row = i + 2
        ck = ws4.cell(row=row, column=1)
        ck.value = k
        ck.font  = fn(bold=True, color=MUTED)
        ck.fill  = F(W if i%2==0 else G1)
        ck.alignment = al("left", "top")
        ck.border = brd()
        cv = ws4.cell(row=row, column=2)
        cv.value = str(v)
        cv.font  = fn(color=TEXT)
        cv.fill  = F(W if i%2==0 else G1)
        cv.alignment = al("left", "top", True)
        cv.border = brd()
        ws4.row_dimensions[row].height = 50 if len(str(v))>60 else 20

    ws4.sheet_properties.tabColor = MUTED

    # ── Guardar ────────────────────────────────────────────────────────────────
    out_path = Path(config.get("output_path", "/mnt/user-data/outputs/"))
    out_path.mkdir(parents=True, exist_ok=True)
    safe_name = re.sub(r'[^a-zA-Z0-9_-]', '_', client_name.lower())
    filename  = f"viral_guiones_{safe_name}_{now.strftime('%Y%m%d')}.xlsx"
    full_path = out_path / filename
    wb.save(str(full_path))
    return str(full_path)

# ─── MAIN ─────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="ECSSTUDIO Viral Scraper")
    parser.add_argument("--config", required=True, help="Ruta al archivo JSON de configuración")
    args = parser.parse_args()

    # Cargar config
    config_path = Path(args.config)
    if not config_path.exists():
        log(f"[ERROR] Config no encontrada: {config_path}")
        sys.exit(1)
    with open(config_path) as f:
        config = json.load(f)

    token       = config.get("apify_token", "")
    min_views   = config.get("min_views", 1_000_000)
    fallback_v  = config.get("fallback_views", 500_000)
    day_range   = config.get("date_range_days", 90)
    results_per = config.get("results_per_hashtag", 100)
    markets     = config.get("markets", [])
    out_lang    = config.get("output_language", "es")

    if not token:
        log("[ERROR] 'apify_token' no definido en el config.")
        sys.exit(1)
    if not markets:
        log("[ERROR] 'markets' vacío en el config.")
        sys.exit(1)

    now    = datetime.now(tz=timezone.utc)
    cutoff = now - timedelta(days=day_range)

    log(f"=== ECSSTUDIO Viral Scraper ===")
    log(f"Cliente: {config.get('client_name','?')} | Nicho: {config.get('client_niche','?')}")
    log(f"Umbral: {min_views:,} views | Ventana: {day_range} días | Output: {out_lang.upper()}")

    # ── Iniciar todos los runs ─────────────────────────────────────────────────
    jobs = []
    for mkt in markets:
        for tag in mkt.get("hashtags", []):
            jobs.append({"market": mkt, "hashtag": tag, "run_id": None, "status": "PENDING"})

    log(f"Lanzando {len(jobs)} trabajos...")
    for i, job in enumerate(jobs):
        run_id = apify_start_run(token, job["hashtag"], results_per)
        if run_id:
            job["run_id"] = run_id
            job["status"] = "RUNNING"
            log(f"  ▶ [{job['market']['id']}] #{job['hashtag']} → {run_id[:8]}...")
        else:
            job["status"] = "FAILED"
        if i < len(jobs) - 1:
            time.sleep(0.7)

    # ── Polling ────────────────────────────────────────────────────────────────
    log("Esperando resultados...")
    done_ids = set()
    all_raw  = []

    while True:
        running = [j for j in jobs if j["status"] == "RUNNING" and j["run_id"] not in done_ids]
        if not running:
            break
        for job in running:
            status = apify_poll(token, job["run_id"])
            if status == "SUCCEEDED":
                done_ids.add(job["run_id"])
                items = apify_get_items(token, job["run_id"])
                for it in items:
                    it["_market"] = job["market"]
                all_raw.extend(items)
                job["status"] = "SUCCEEDED"
                log(f"  ✓ [{job['market']['id']}] #{job['hashtag']} → {len(items)} items")
            elif status in ("FAILED", "TIMED-OUT", "ABORTED"):
                done_ids.add(job["run_id"])
                job["status"] = "FAILED"
                log(f"  ✗ [{job['market']['id']}] #{job['hashtag']} → {status}")
        if any(j["status"] == "RUNNING" for j in jobs):
            time.sleep(POLL_SECS)

    log(f"Total bruto: {len(all_raw)} items")

    # ── Deduplicar ────────────────────────────────────────────────────────────
    seen = set()
    unique = []
    for item in all_raw:
        vid_id = item.get("id", "")
        if vid_id and vid_id not in seen:
            seen.add(vid_id)
            unique.append(item)
    log(f"Únicos: {len(unique)}")

    # ── Filtrar y procesar ────────────────────────────────────────────────────
    filtered = [item for item in unique if item.get("playCount", 0) >= min_views]
    log(f"Con {min_views:,}+ views: {len(filtered)}")

    # Fallback si hay muy pocos
    if len(filtered) < 5:
        log(f"⚠ Menos de 5 resultados con {min_views:,}+ views. Aplicando fallback a {fallback_v:,}...")
        filtered = [item for item in unique if item.get("playCount", 0) >= fallback_v]
        log(f"  Con {fallback_v:,}+ views: {len(filtered)}")

    filtered.sort(key=lambda x: x.get("playCount", 0), reverse=True)

    log("Procesando transcripciones y traducciones...")
    results = []
    for i, item in enumerate(filtered):
        mkt = item.get("_market", markets[0])
        r   = process_item(item, mkt, out_lang, cutoff)
        if r:
            results.append(r)
            tmark = "→ traducido" if r["translated"] else ""
            log(f"  [{i+1}/{len(filtered)}] @{r['handle']} | {r['views_fmt']} | {r['date']} {tmark}")

    log(f"Videos válidos: {len(results)}")

    if not results:
        log("⚠ Sin resultados que cumplan los criterios. Ajustar parámetros e intentar de nuevo.")
        sys.exit(0)

    # ── Generar XLSX ───────────────────────────────────────────────────────────
    log("Generando XLSX...")
    out_file = build_xlsx(results, config, now)
    log(f"✓ Guardado: {out_file}")
    log(f"  Videos: {len(results)} | Traducidos: {sum(1 for r in results if r['translated'])}")

    # Resumen en consola
    print("\n" + "="*55)
    print("RESUMEN")
    print("="*55)
    for r in results[:10]:
        print(f"{r['market_flag']} @{r['handle']} | {r['views_fmt']} | {r['date']} | {r['lang_orig']}")
        print(f"   └ {r['transcript_out'][:100]}...")
    if len(results) > 10:
        print(f"   ... y {len(results)-10} videos más en el XLSX")
    print(f"\nArchivo: {out_file}")

if __name__ == "__main__":
    main()
